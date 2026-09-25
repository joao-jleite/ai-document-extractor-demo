"""Offline tests: validators, business rules and exporters (no API key needed)."""

import json
from datetime import date
from pathlib import Path

import pytest

from app.exporters import write_pdf, write_xlsx
from app.extractor import InvalidFileError, prepare_input, sniff_media_type
from app.schema import ExtractedDocument
from app.validation import (
    is_valid_cnpj, is_valid_cpf, is_valid_rut, nfe_key_dv_ok, validate_document,
)

ROOT = Path(__file__).resolve().parent.parent
TRUTH = ROOT / "samples" / "truth"
TODAY = date(2026, 9, 25)


def load_truth(stem: str) -> ExtractedDocument:
    data = json.loads((TRUTH / f"{stem}.json").read_text(encoding="utf-8"))
    return ExtractedDocument.model_validate({k: v for k, v in data.items() if not k.startswith("_")})


@pytest.mark.parametrize("value,ok", [
    ("11.222.333/0001-81", True),
    ("11222333000181", True),
    ("11.222.333/0001-82", False),
    ("12.ABC.345/01DE-35", True),   # alphanumeric CNPJ (2026 format)
    ("12.ABC.345/01DE-36", False),
    ("00.000.000/0000-00", False),
])
def test_cnpj(value, ok):
    assert is_valid_cnpj(value) is ok


@pytest.mark.parametrize("value,ok", [
    ("11.111.111-1", True), ("11111111-1", True), ("11.111.111-2", False),
    ("77.123.456-9", True), ("96.555.444-0", True), ("12.345.678-5", True), ("12.345.678-K", False),
])
def test_rut(value, ok):
    assert is_valid_rut(value) is ok


def test_cpf():
    assert is_valid_cpf("529.982.247-25")
    assert not is_valid_cpf("529.982.247-24")
    assert not is_valid_cpf("111.111.111-11")


def test_nfe_key_from_sample():
    key = load_truth("danfe-exemplo-industrial").nfe_access_key
    assert nfe_key_dv_ok(key)
    assert not nfe_key_dv_ok(key[:-1] + str((int(key[-1]) + 1) % 10))


def test_clean_danfe_passes_everything():
    report = validate_document(load_truth("danfe-exemplo-industrial"), today=TODAY)
    assert report.status == "ok", [c for c in report.checks if c.status != "ok"]
    codes = {c.code for c in report.checks}
    assert {"required_fields", "supplier_tax_id", "buyer_tax_id", "line_math",
            "items_vs_subtotal", "grand_total", "nfe_key", "dates", "currency"} <= codes


def test_purchase_order_passes_with_cnpj_and_rut():
    report = validate_document(load_truth("orden-compra-andina"), today=TODAY)
    assert report.status == "ok"
    titles = {c.title for c in report.checks}
    assert "Supplier CNPJ" in titles and "Buyer RUT" in titles


def test_photo_sample_flags_the_deliberate_typo():
    report = validate_document(load_truth("foto-danfe-parafusos"), today=TODAY)
    assert report.status == "review"
    warnings = {c.code: c.detail for c in report.checks if c.status == "warning"}
    assert set(warnings) == {"line_math", "items_vs_subtotal"}
    assert "472.50" in warnings["line_math"] and "prints 427.50" in warnings["line_math"]
    assert "-45.00" in warnings["items_vs_subtotal"]


def test_missing_and_invalid_fields_are_errors():
    doc = load_truth("orden-compra-andina").model_copy(deep=True)
    doc.buyer.tax_id = "11.111.111-9"
    doc.grand_total = None
    doc.issue_date = "2026-02-31"
    report = validate_document(doc, today=TODAY)
    errors = {c.code for c in report.checks if c.status == "error"}
    assert {"buyer_tax_id", "required_fields", "dates"} <= errors


def test_clp_decimals_warning():
    doc = load_truth("orden-compra-andina").model_copy(deep=True)
    doc.currency = "CLP"
    report = validate_document(doc, today=TODAY)
    assert any(c.code == "currency_decimals" for c in report.checks)


def test_file_sniffing_and_rejection(tmp_path):
    assert sniff_media_type((ROOT / "samples" / "orden-compra-andina.pdf").read_bytes()) == "application/pdf"
    assert sniff_media_type((ROOT / "samples" / "foto-danfe-parafusos.jpg").read_bytes()) == "image/jpeg"
    with pytest.raises(InvalidFileError):
        prepare_input(b"hello world, not a pdf", "fake.pdf")
    with pytest.raises(InvalidFileError):
        prepare_input(b"", "empty.png")
    with pytest.raises(InvalidFileError):
        prepare_input(b"%PDF-1.4 truncated", "cut.pdf")
    prepared = prepare_input((ROOT / "samples" / "foto-danfe-parafusos.jpg").read_bytes(), "photo.jpg")
    assert prepared.kind == "image" and prepared.media_type == "image/jpeg"


def test_exporters_write_files(tmp_path):
    doc = load_truth("foto-danfe-parafusos")
    report = validate_document(doc, today=TODAY)
    meta = {"filename": "foto.jpg", "model": "test", "generated_at": "2026-09-25", "seconds": 1, "attempts": 1}
    xlsx = write_xlsx(tmp_path / "out.xlsx", doc, report, meta)
    pdf = write_pdf(tmp_path / "out.pdf", doc, report, meta)
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Header", "Items", "Validation"]
    assert "DEMO" in wb["Header"]["A1"].value
    assert pdf.read_bytes()[:5] == b"%PDF-"


def _key_with_dv(first43: str) -> str:
    """Append the mod-11 check digit (letters count as ord(c) - 48, like the alphanumeric CNPJ)."""
    total, weight = 0, 2
    for ch in reversed(first43):
        total += (ord(ch) - 48) * weight
        weight = 2 if weight == 9 else weight + 1
    rest = total % 11
    return first43 + str(0 if rest < 2 else 11 - rest)


def test_nfe_key_with_alphanumeric_cnpj_issuer():
    # Fictitious key: SP, 2026-09, issuer = the alphanumeric CNPJ format example, NF 4217.
    key = _key_with_dv("352609" + "12ABC34501DE35" + "55" + "001" + "000004217" + "1" + "73920184")
    assert len(key) == 44 and nfe_key_dv_ok(key)
    assert not nfe_key_dv_ok(key[:-1] + str((int(key[-1]) + 1) % 10))
    assert not nfe_key_dv_ok(key.replace("12ABC", "12ABD"))  # one letter changed -> DV no longer matches

    doc = load_truth("danfe-exemplo-industrial").model_copy(deep=True)
    doc.supplier.tax_id = "12.ABC.345/01DE-35"
    # As the model copies it: 4-character groups separated by spaces, compacted by the schema.
    doc = ExtractedDocument.model_validate({**doc.model_dump(), "nfe_access_key": " ".join(
        key[i:i + 4] for i in range(0, 44, 4))})
    assert doc.nfe_access_key == key
    check = next(c for c in validate_document(doc, today=TODAY).checks if c.code == "nfe_key")
    assert check.status == "ok", check.detail


def test_pdf_title_with_markup_characters(tmp_path):
    doc = load_truth("orden-compra-andina").model_copy(deep=True)
    doc.document_number = "OC <7> & 8"  # would break reportlab's paragraph markup if not escaped
    report = validate_document(doc, today=TODAY)
    meta = {"filename": "po.pdf", "model": "test", "generated_at": "2026-09-25", "seconds": 1, "attempts": 1}
    assert write_pdf(tmp_path / "po.pdf", doc, report, meta).read_bytes()[:5] == b"%PDF-"


def test_user_documents_are_not_labelled_fictitious(tmp_path):
    from openpyxl import load_workbook

    from app.exporters import DEMO_BANNER, DEMO_BANNER_OWN

    doc = load_truth("orden-compra-andina")
    report = validate_document(doc, today=TODAY)
    base = {"filename": "po.pdf", "model": "test", "generated_at": "2026-09-25", "seconds": 1, "attempts": 1}
    own = write_xlsx(tmp_path / "own.xlsx", doc, report, {**base, "sample": False})
    sample = write_xlsx(tmp_path / "sample.xlsx", doc, report, {**base, "sample": True})
    assert load_workbook(own)["Header"]["A1"].value == DEMO_BANNER_OWN
    assert "fictitious" not in DEMO_BANNER_OWN and "DEMO" in DEMO_BANNER_OWN
    assert load_workbook(sample)["Header"]["A1"].value == DEMO_BANNER
    assert write_pdf(tmp_path / "own.pdf", doc, report, {**base, "sample": False}).is_file()
