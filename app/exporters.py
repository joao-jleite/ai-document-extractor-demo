"""XLSX (openpyxl) and PDF (reportlab) outputs. Both are clearly marked DEMO."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from .schema import ExtractedDocument, ValidationReport
from .validation import CURRENCY_DECIMALS, line_mismatches

DEMO_BANNER = "DEMO  -  fictitious data  -  AI Document Extractor"
# A user's own document is not fictitious: its outputs are only marked DEMO.
DEMO_BANNER_OWN = "DEMO  -  AI Document Extractor  -  review before use"


def _is_sample(meta: dict) -> bool:
    """True for the bundled fictitious samples (pipeline matches them by SHA-256)."""
    return bool(meta.get("sample", True))
DOC_TYPE_LABEL = {
    "nfe_danfe": "NF-e (DANFE)",
    "purchase_order": "Purchase order",
    "commercial_invoice": "Commercial invoice",
    "other": "Other document",
}
STATUS_LABEL = {"ok": "OK", "warning": "WARNING", "error": "ERROR", "info": "INFO"}


def excel_money_format(currency: str | None, unit_price: bool = False) -> str:
    """Excel number format with the currency symbol. Unit prices may carry up
    to 4 decimals (common on NF-e), so they get optional extra digits."""
    dec = CURRENCY_DECIMALS.get(currency or "", 2)
    body = "#,##0" + ("." + "0" * dec if dec else "")
    if unit_price and dec:
        body += "##"
    prefix = {"BRL": '"R$ "', "USD": '"US$ "', "CLP": '"CLP$ "', "EUR": '"€ "'}.get(currency or "")
    return f"{prefix}{body}" if prefix else f'{body} "{currency or ""}"'


def money(x: float | None, currency: str | None) -> str:
    if x is None:
        return "-"
    dec = CURRENCY_DECIMALS.get(currency or "", 2)
    return f"{x:,.{dec}f}"


def unit_price(x: float | None, currency: str | None) -> str:
    """Unit prices keep at least the currency decimals and up to 4 (NF-e style)."""
    if x is None:
        return "-"
    dec = max(CURRENCY_DECIMALS.get(currency or "", 2), 2)
    text = f"{x:,.4f}"
    head, tail = text.split(".")
    tail = tail[:dec] + tail[dec:].rstrip("0")
    return f"{head}.{tail}"


def _as_date(value: str | None):
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return value  # keep the raw text; validation already flagged it


def header_rows(doc: ExtractedDocument, report: ValidationReport) -> list[tuple[str, object, str]]:
    """(label, value, kind) rows shared by XLSX and PDF. kind: text | money | date."""
    s, b = doc.supplier, doc.buyer
    return [
        ("Document type", DOC_TYPE_LABEL.get(doc.document_type, doc.document_type), "text"),
        ("Document number", doc.document_number, "text"),
        ("Series", doc.series, "text"),
        ("Issue date", _as_date(doc.issue_date), "date"),
        ("Due / delivery date", _as_date(doc.due_or_delivery_date), "date"),
        ("Currency", doc.currency, "text"),
        ("Language", doc.document_language, "text"),
        ("Supplier", s.name, "text"),
        ("Supplier tax ID", f"{s.tax_id or '-'} ({s.tax_id_type or '?'})", "text"),
        ("Supplier country", s.country, "text"),
        ("Supplier address", s.address, "text"),
        ("Buyer", b.name, "text"),
        ("Buyer tax ID", f"{b.tax_id or '-'} ({b.tax_id_type or '?'})", "text"),
        ("Buyer country", b.country, "text"),
        ("Buyer address", b.address, "text"),
        ("Items subtotal", doc.items_subtotal, "money"),
        ("Discount", doc.discount, "money"),
        ("Freight", doc.freight, "money"),
        ("Insurance", doc.insurance, "money"),
        ("Other charges", doc.other_charges, "money"),
        ("Taxes added", doc.tax_added, "money"),
        ("Grand total", doc.grand_total, "money"),
        ("Payment terms", doc.payment_terms, "text"),
        ("NF-e access key", doc.nfe_access_key, "text"),
        ("Notes", doc.notes, "text"),
        ("Validation", "All checks passed" if report.status == "ok" else "Needs review", "text"),
    ]


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

_BANNER_FILL = PatternFill("solid", fgColor="FDE68A")
_HEAD_FILL = PatternFill("solid", fgColor="1E293B")
_ZEBRA_FILL = PatternFill("solid", fgColor="F8FAFC")
_FLAG_FILL = PatternFill("solid", fgColor="FEF3C7")
_STATUS_FILL = {
    "ok": PatternFill("solid", fgColor="DCFCE7"),
    "warning": PatternFill("solid", fgColor="FEF3C7"),
    "error": PatternFill("solid", fgColor="FEE2E2"),
    "info": PatternFill("solid", fgColor="E0F2FE"),
}
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(bottom=_THIN)


def _banner(ws, width: int, subtitle: str, text: str = DEMO_BANNER) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value=text)
    c.fill, c.font = _BANNER_FILL, Font(bold=True, size=12, color="78350F")
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font = Font(italic=True, size=9, color="64748B")


def _header_row(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.fill, c.font = _HEAD_FILL, Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 18


def write_xlsx(path: Path, doc: ExtractedDocument, report: ValidationReport, meta: dict) -> Path:
    wb = Workbook()
    cur = doc.currency
    money_fmt = excel_money_format(cur)
    # Offline runs did not call Claude: say "Extraction: ..." instead of naming a model.
    who = "Extraction" if meta.get("mode") == "offline" else "Model"
    subtitle = (
        f"Source: {meta.get('filename')}  |  {who}: {meta.get('model')}  |  "
        f"Generated: {meta.get('generated_at')}"
    )
    banner = DEMO_BANNER if _is_sample(meta) else DEMO_BANNER_OWN

    # --- Header -----------------------------------------------------------
    ws = wb.active
    ws.title = "Header"
    _banner(ws, 2, subtitle, banner)
    _header_row(ws, 4, ["Field", "Value"])
    for i, (label, value, kind) in enumerate(header_rows(doc, report), start=5):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True, color="334155")
        c = ws.cell(row=i, column=2, value=value)
        if kind == "money" and value is not None:
            c.number_format = money_fmt
            c.alignment = Alignment(horizontal="left")
        elif kind == "date" and isinstance(value, date):
            c.number_format = "yyyy-mm-dd"
            c.alignment = Alignment(horizontal="left")
        else:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if label == "Validation":
            c.fill = _STATUS_FILL["ok" if report.status == "ok" else "warning"]
            c.font = Font(bold=True)
        for col in (1, 2):
            ws.cell(row=i, column=col).border = _BORDER
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "A5"

    # --- Items ------------------------------------------------------------
    wi = wb.create_sheet("Items")
    heads = ["#", "Code", "Description", "Qty", "Unit", "Unit price", "Line total", "Qty x price", "Check"]
    _banner(wi, len(heads), subtitle, banner)
    _header_row(wi, 4, heads)
    flagged = line_mismatches(doc)
    dec = CURRENCY_DECIMALS.get(cur or "", 2)
    row = 5
    for idx, item in enumerate(doc.items):
        expected = (
            round(item.quantity * item.unit_price, dec)
            if item.quantity is not None and item.unit_price is not None else None
        )
        values = [item.line_number or idx + 1, item.code, item.description, item.quantity, item.unit,
                  item.unit_price, item.line_total, expected, "MISMATCH" if idx in flagged else "ok"]
        for col, v in enumerate(values, start=1):
            c = wi.cell(row=row, column=col, value=v)
            c.border = _BORDER
            if col in (7, 8):
                c.number_format = money_fmt
            if col == 6:
                c.number_format = excel_money_format(cur, unit_price=True)
            if col == 4:
                c.number_format = "#,##0.###"
            if idx in flagged:
                c.fill = _FLAG_FILL
            elif idx % 2:
                c.fill = _ZEBRA_FILL
        wi.cell(row=row, column=9).font = Font(bold=idx in flagged, color="B45309" if idx in flagged else "15803D")
        row += 1

    # Totals block (values, not formulas, so every viewer shows the same numbers)
    row += 1
    items_sum = report.computed.get("items_sum")
    totals = [
        ("Sum of line totals (computed)", items_sum),
        ("Items subtotal (printed)", doc.items_subtotal),
        ("Difference", None if items_sum is None or doc.items_subtotal is None
         else round(items_sum - doc.items_subtotal, dec)),
        ("Grand total (printed)", doc.grand_total),
        ("Grand total (expected from components)", report.computed.get("expected_total")),
    ]
    for label, value in totals:
        # Label spans C:F so long captions do not wrap in the narrow price column
        wi.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        lab = wi.cell(row=row, column=3, value=label)
        lab.font = Font(bold=True, color="334155")
        lab.alignment = Alignment(horizontal="right")
        c = wi.cell(row=row, column=7, value=value)
        c.number_format = money_fmt
        c.font = Font(bold=True)
        if label == "Difference" and value not in (None, 0):
            c.fill = _FLAG_FILL
        row += 1
    widths = [5, 14, 46, 9, 7, 14, 16, 16, 11]
    for col, w in enumerate(widths, start=1):
        wi.column_dimensions[get_column_letter(col)].width = w
    wi.freeze_panes = "A5"
    wi.auto_filter.ref = f"A4:I{4 + len(doc.items)}"

    # --- Validation ---------------------------------------------------------
    wv = wb.create_sheet("Validation")
    _banner(wv, 3, subtitle, banner)
    _header_row(wv, 4, ["Status", "Check", "Detail"])
    order = {"error": 0, "warning": 1, "ok": 2, "info": 3}
    for i, chk in enumerate(sorted(report.checks, key=lambda c: order[c.status]), start=5):
        s = wv.cell(row=i, column=1, value=STATUS_LABEL[chk.status])
        s.fill, s.font = _STATUS_FILL[chk.status], Font(bold=True)
        wv.cell(row=i, column=2, value=chk.title).font = Font(bold=True)
        d = wv.cell(row=i, column=3, value=chk.detail)
        d.alignment = Alignment(wrap_text=True, vertical="top")
        for col in (1, 2, 3):
            wv.cell(row=i, column=col).border = _BORDER
            wv.cell(row=i, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    wv.column_dimensions["A"].width = 11
    wv.column_dimensions["B"].width = 22
    wv.column_dimensions["C"].width = 95
    wv.freeze_panes = "A5"

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    wb.properties.title = "DEMO - AI Document Extractor output"
    wb.properties.creator = "AI Document Extractor (demo)"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# PDF report
# ---------------------------------------------------------------------------

INK = colors.HexColor("#0F172A")
MUTED = colors.HexColor("#64748B")
LINE = colors.HexColor("#E2E8F0")
PDF_STATUS = {
    "ok": (colors.HexColor("#DCFCE7"), colors.HexColor("#166534")),
    "warning": (colors.HexColor("#FEF3C7"), colors.HexColor("#92400E")),
    "error": (colors.HexColor("#FEE2E2"), colors.HexColor("#991B1B")),
    "info": (colors.HexColor("#E0F2FE"), colors.HexColor("#075985")),
}


def _page_decor(canvas, doc_template, sample: bool = True):
    """Header strip, diagonal DEMO watermark and footer on every page.
    "Fictitious" is printed only for the bundled samples (not for a user's own file)."""
    w, h = A4
    canvas.saveState()
    # watermark
    canvas.setFillColor(colors.Color(0.85, 0.1, 0.1, alpha=0.07))
    canvas.setFont("Helvetica-Bold", 120)
    canvas.translate(w / 2, h / 2)
    canvas.rotate(35)
    canvas.drawCentredString(0, -40, "DEMO")
    canvas.restoreState()

    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#FDE68A"))
    canvas.rect(0, h - 12 * mm, w, 12 * mm, stroke=0, fill=1)
    canvas.setFillColor(colors.HexColor("#78350F"))
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawString(15 * mm, h - 7.5 * mm, "DEMO  -  AI Document Extractor  -  extraction report")
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(w - 15 * mm, h - 7.5 * mm,
                           "All data is fictitious" if sample else "Demo output - review before use")
    canvas.setStrokeColor(LINE)
    canvas.line(15 * mm, 14 * mm, w - 15 * mm, 14 * mm)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 7.5)
    canvas.drawString(15 * mm, 9 * mm, "DEMO - fictitious data. Generated automatically; review before use."
                      if sample else "DEMO - generated automatically by a demo tool; review before use.")
    canvas.drawRightString(w - 15 * mm, 9 * mm, f"Page {doc_template.page}")
    canvas.restoreState()


def write_pdf(path: Path, doc: ExtractedDocument, report: ValidationReport, meta: dict) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=styles["Normal"], fontName="Helvetica", fontSize=8.5,
                          leading=11, textColor=INK)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, leading=9.5, textColor=MUTED)
    label = ParagraphStyle("label", parent=base, fontName="Helvetica-Bold", textColor=colors.HexColor("#334155"))
    right = ParagraphStyle("right", parent=base, alignment=TA_RIGHT)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=21)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=10, spaceAfter=4)
    cur = doc.currency
    P = lambda text, st=base: Paragraph(str(text if text not in (None, "") else "-")  # noqa: E731
                                        .replace("&", "&amp;").replace("<", "&lt;"), st)

    story = []
    title = f"{DOC_TYPE_LABEL.get(doc.document_type, 'Document')} {doc.document_number or ''}".strip()
    story += [P(title, h1), Spacer(1, 2 * mm)]  # P() escapes '&' and '<' (reportlab markup)
    if meta.get("mode") == "offline":  # no model call: no model name, time or attempts to report
        run_line = (f"Source file: {meta.get('filename')}  |  Extraction: {meta.get('model')}  |  "
                    f"Processed: {meta.get('generated_at')}")
    elif meta.get("mode") == "replay":  # no call now: the time is the saved live run's, attempts do not apply
        run_line = (f"Source file: {meta.get('filename')}  |  Model: {meta.get('model')}  |  "
                    f"Processed: {meta.get('generated_at')}  |  Extraction time of the saved live run: "
                    f"{meta.get('seconds')} s")
    else:
        run_line = (f"Source file: {meta.get('filename')}  |  Model: {meta.get('model')}  |  "
                    f"Processed: {meta.get('generated_at')}  |  Extraction time: {meta.get('seconds')} s  |  "
                    f"Attempts: {meta.get('attempts')}")
    story.append(P(run_line, small))
    story.append(Spacer(1, 4 * mm))

    # Status box
    n_warn, n_err = report.counts.get("warning", 0), report.counts.get("error", 0)
    key = "ok" if report.status == "ok" else ("error" if n_err else "warning")
    bg, fg = PDF_STATUS[key]
    parts = [f"{n} {word}{'s' if n != 1 else ''}" for n, word in ((n_err, "error"), (n_warn, "warning")) if n]
    status_text = (
        f"All {report.counts.get('ok', 0)} checks passed" if report.status == "ok"
        else f"Needs review: {', '.join(parts)} - {report.counts.get('ok', 0)} checks passed"
    )
    box = Table([[Paragraph(f"<b>{status_text}</b>", ParagraphStyle("st", parent=base, textColor=fg, fontSize=10))]],
                colWidths=[180 * mm])
    box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), bg), ("BOX", (0, 0), (-1, -1), 0.6, fg),
                             ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story.append(box)

    # Parties side by side
    story.append(Paragraph("Parties", h2))
    def party_cell(role, p):
        cell = [P(role, label), P(p.name), P(f"{p.tax_id or '-'} ({p.tax_id_type or '?'})", small)]
        cell += [P(v, small) for v in (p.address, p.country) if v]  # skip empty lines
        return cell
    parties = Table([[party_cell("Supplier", doc.supplier), party_cell("Buyer", doc.buyer)]],
                    colWidths=[90 * mm, 90 * mm])
    parties.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (0, 0), (0, 0), 0.5, LINE),
                                 ("BOX", (1, 0), (1, 0), 0.5, LINE), ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                 ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(parties)

    # Header fields (2 x 2 grid of label/value)
    story.append(Paragraph("Document", h2))
    fields = [
        ("Number", doc.document_number), ("Series", doc.series),
        ("Issue date", doc.issue_date), ("Due / delivery", doc.due_or_delivery_date),
        ("Currency", cur), ("Payment terms", doc.payment_terms),
        ("NF-e access key", doc.nfe_access_key), ("Language", doc.document_language),
    ]
    grid = [[P(fields[i][0], label), P(fields[i][1]), P(fields[i + 1][0], label), P(fields[i + 1][1])]
            for i in range(0, len(fields), 2)]
    t = Table(grid, colWidths=[28 * mm, 62 * mm, 28 * mm, 62 * mm])
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.4, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t)

    # Items
    story.append(Paragraph(f"Items ({len(doc.items)})", h2))
    flagged = line_mismatches(doc)
    rows = [[P(h, label) for h in ("#", "Code", "Description", "Qty", "Unit", "Unit price", "Line total")]]
    for idx, it in enumerate(doc.items):
        rows.append([P(it.line_number or idx + 1), P(it.code), P(it.description),
                     P(f"{it.quantity:g}" if it.quantity is not None else "-", right), P(it.unit),
                     P(unit_price(it.unit_price, cur), right),
                     P(money(it.line_total, cur), right)])
    items = Table(rows, colWidths=[8 * mm, 20 * mm, 72 * mm, 14 * mm, 12 * mm, 24 * mm, 30 * mm], repeatRows=1)
    ist = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
           ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP")]
    for idx in flagged:
        ist.append(("BACKGROUND", (0, idx + 1), (-1, idx + 1), PDF_STATUS["warning"][0]))
    items.setStyle(TableStyle(ist))
    story.append(items)

    # Totals
    tot_rows = [
        ("Sum of line totals (computed)", report.computed.get("items_sum")),
        ("Items subtotal (printed)", doc.items_subtotal),
        ("Discount", doc.discount), ("Freight", doc.freight), ("Insurance", doc.insurance),
        ("Other charges", doc.other_charges), ("Taxes added", doc.tax_added),
        ("Grand total (printed)", doc.grand_total),
    ]
    tot = Table([[P(l, right), P(f"{money(v, cur)} {cur or ''}", right)] for l, v in tot_rows if v is not None],
                colWidths=[140 * mm, 40 * mm])
    tot.setStyle(TableStyle([("LINEABOVE", (0, -1), (-1, -1), 0.8, INK),
                             ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")]))
    story += [Spacer(1, 2 * mm), tot]

    # Validation checklist
    order = {"error": 0, "warning": 1, "ok": 2, "info": 3}
    vrows = []
    vstyle = [("LINEBELOW", (0, 0), (-1, -1), 0.4, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP")]
    for r, chk in enumerate(sorted(report.checks, key=lambda c: order[c.status])):
        bg, fg = PDF_STATUS[chk.status]
        vrows.append([Paragraph(f"<b>{STATUS_LABEL[chk.status]}</b>",
                                ParagraphStyle("s", parent=base, textColor=fg, fontSize=7.5)),
                      P(chk.title, label), P(chk.detail)])
        vstyle.append(("BACKGROUND", (0, r), (0, r), bg))
    vt = Table(vrows, colWidths=[20 * mm, 38 * mm, 122 * mm])
    vt.setStyle(TableStyle(vstyle))
    story.append(KeepTogether([Paragraph("Validation", h2), vt]))

    pdf = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=20 * mm,
        bottomMargin=20 * mm, title="DEMO - Extraction report", author="AI Document Extractor (demo)",
    )
    decor = lambda canvas, tpl: _page_decor(canvas, tpl, _is_sample(meta))  # noqa: E731
    pdf.build(story, onFirstPage=decor, onLaterPages=decor)
    return path
