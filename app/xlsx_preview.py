"""Render a generated .xlsx as HTML, reading the real file with openpyxl.

This is what the "Preview XLSX" link shows: values, number formats, fills,
fonts, merged cells and column widths come from the workbook itself, so the
preview matches what Excel/LibreOffice will open.
"""

from __future__ import annotations

import html
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _format_number(value: float, fmt: str) -> str:
    """Tiny subset of Excel number formats used by our exporter:
    optional quoted prefix/suffix + #,##0 with fixed and optional decimals."""
    prefix = "".join(re.findall(r'^"([^"]*)"', fmt))
    suffix = "".join(re.findall(r'"([^"]*)"$', fmt)) if not prefix else ""
    core = re.sub(r'"[^"]*"', "", fmt).strip()
    fixed = optional = 0
    if "." in core:
        decimals = core.split(".", 1)[1]
        fixed, optional = decimals.count("0"), decimals.count("#")
    text = f"{value:,.{fixed + optional}f}"
    if optional:  # trim optional (#) decimals, never below the fixed ones
        int_part, dec_part = text.split(".")
        dec_part = dec_part[:fixed] + dec_part[fixed:].rstrip("0")
        text = int_part + ("." + dec_part if dec_part else "")
    return f"{prefix}{text}{(' ' + suffix) if suffix else ''}".strip()


def _cell_text(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        fmt = cell.number_format or "General"
        return str(v) if fmt == "General" else _format_number(float(v), fmt)
    return str(v)


def _style(cell) -> str:
    css = []
    fill = cell.fill
    if fill is not None and fill.fill_type == "solid" and fill.fgColor is not None and fill.fgColor.rgb:
        rgb = str(fill.fgColor.rgb)[-6:]
        if re.fullmatch(r"[0-9A-Fa-f]{6}", rgb):
            css.append(f"background:#{rgb}")
    font = cell.font
    if font is not None:
        if font.b:
            css.append("font-weight:600")
        if font.i:
            css.append("font-style:italic")
        if font.color is not None and font.color.rgb and isinstance(font.color.rgb, str):
            css.append(f"color:#{font.color.rgb[-6:]}")
        if font.sz:
            css.append(f"font-size:{float(font.sz) * 1.25:.1f}px")
    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
        if cell.alignment is None or cell.alignment.horizontal in (None, "general"):
            css.append("text-align:right")
    if cell.alignment is not None and cell.alignment.horizontal in ("left", "right", "center"):
        css.append(f"text-align:{cell.alignment.horizontal}")
    return ";".join(css)


def render_xlsx_html(path: Path, title: str = "XLSX preview") -> str:
    wb = load_workbook(path)
    tabs, sheets = [], []
    for i, ws in enumerate(wb.worksheets):
        merged = {}
        skip = set()
        for rng in ws.merged_cells.ranges:
            merged[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        skip.add((r, c))
        # Pad with empty columns/rows so the preview reads like a spreadsheet grid.
        cols = max(ws.max_column, 12)
        rows = max(ws.max_row, 40)
        widths = [(ws.column_dimensions[get_column_letter(c)].width or 9) * 7.2 for c in range(1, cols + 1)]
        colgroup = "<col class='rh'>" + "".join(f"<col style='width:{w:.0f}px'>" for w in widths)
        head = "<tr><th class='corner'></th>" + "".join(
            f"<th>{get_column_letter(c)}</th>" for c in range(1, cols + 1)) + "</tr>"
        body = []
        for r in range(1, rows + 1):
            cells = [f"<th class='rownum'>{r}</th>"]
            for c in range(1, cols + 1):
                if (r, c) in skip:
                    continue
                cell = ws.cell(row=r, column=c)
                span = merged.get((r, c))
                attrs = f" rowspan='{span[0]}' colspan='{span[1]}'" if span else ""
                cells.append(f"<td{attrs} style='{_style(cell)}'>{html.escape(_cell_text(cell))}</td>")
            body.append("<tr>" + "".join(cells) + "</tr>")
        active = " active" if i == 0 else ""
        tabs.append(f"<button class='tab{active}' data-i='{i}'>{html.escape(ws.title)}</button>")
        sheets.append(f"<div class='sheet{active}' id='s{i}'><table><colgroup>{colgroup}</colgroup>"
                      f"<thead>{head}</thead><tbody>{''.join(body)}</tbody></table></div>")

    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title>
<style>
:root{{--grid:#E2E8F0;--head:#F1F5F9;--ink:#0F172A}}
*{{box-sizing:border-box}} body{{margin:0;font:13px/1.35 "Segoe UI",Calibri,Arial,sans-serif;color:var(--ink);background:#fff}}
.bar{{display:flex;align-items:center;gap:12px;padding:10px 16px;background:#107C41;color:#fff}}
.bar b{{font-size:14px}} .bar span{{opacity:.85;font-size:12px}}
.pill{{background:#FDE68A;color:#78350F;font-weight:700;font-size:11px;padding:2px 8px;border-radius:999px}}
.wrap{{overflow:auto;height:calc(100vh - 86px)}}
.sheet{{display:none}} .sheet.active{{display:block}}
table{{border-collapse:collapse;table-layout:fixed}}
th,td{{border:1px solid var(--grid);padding:3px 6px;white-space:pre-wrap;vertical-align:top;overflow:hidden}}
thead th{{background:var(--head);font-weight:500;color:#475569;position:sticky;top:0;font-size:11px;text-align:center}}
th.rownum{{background:var(--head);color:#475569;font-weight:500;font-size:11px;text-align:center;width:36px}}
col.rh{{width:36px}}
.tabs{{display:flex;gap:2px;padding:6px 8px 0;background:#F8FAFC;border-top:1px solid var(--grid);position:fixed;bottom:0;left:0;right:0;height:40px}}
.tab{{border:1px solid var(--grid);border-bottom:none;background:#EEF2F6;padding:6px 16px;font:inherit;cursor:pointer;border-radius:6px 6px 0 0}}
.tab.active{{background:#fff;color:#107C41;font-weight:600;border-bottom:2px solid #107C41}}
</style></head><body>
<div class="bar"><b>{html.escape(path.name)}</b><span>Rendered from the generated workbook (openpyxl)</span><span class="pill">DEMO</span></div>
<div class="wrap">{''.join(sheets)}</div>
<div class="tabs">{''.join(tabs)}</div>
<script>
document.querySelectorAll('.tab').forEach(b=>b.onclick=()=>{{
  document.querySelectorAll('.tab,.sheet').forEach(e=>e.classList.remove('active'));
  b.classList.add('active'); document.getElementById('s'+b.dataset.i).classList.add('active');
}});
const q=new URLSearchParams(location.search).get('sheet'); if(q){{const b=[...document.querySelectorAll('.tab')].find(t=>t.textContent===q); if(b) b.click();}}
</script></body></html>"""
